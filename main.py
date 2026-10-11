from fastapi import BackgroundTasks, Depends, FastAPI, File, HTTPException, Query, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse
from pathlib import Path
from pydantic import BaseModel, EmailStr, Field
from typing import Any, Optional, List
import csv
import hashlib
import io
import json
import os
import uuid
from datetime import datetime, timezone

from auth import create_access_token, get_current_user, hash_password, public_user, require_admin, verify_password
from llm_gateway import LLMGateway, ProviderError
from mongo_store import get_store, is_production
from pipeline import generate_pipeline
from question_bank import QuestionBank, QuestionStatus
from execution_engine import verify_in_sandbox


def _now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


app = FastAPI(
    title="SQLArena AI Question Generator",
    description="7-step AI pipeline for SQL question generation",
    version="2.0.0"
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=[origin.strip() for origin in os.getenv("CORS_ORIGINS", "http://localhost:5173,http://127.0.0.1:5173").split(",") if origin.strip()],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


class ApiPrefixMiddleware:
    """Strip exactly one leading `/api` segment.

    This is the single place the prefix is removed. `vercel.json` rewrites
    `/api/*` here with the prefix intact, and the Vite dev proxy does the same,
    so local and deployed routing match. `/apifoo` is left alone.
    """

    def __init__(self, app):
        self.app = app

    async def __call__(self, scope, receive, send):
        if scope["type"] == "http" and (scope["path"] == "/api" or scope["path"].startswith("/api/")):
            scope = dict(scope)
            scope["path"] = scope["path"][4:] or "/"
        await self.app(scope, receive, send)


app.add_middleware(ApiPrefixMiddleware)

_STATIC_DIR = Path(__file__).parent / "static"
_FRONTEND_DIR = Path(__file__).parent / "dist"
_LEGACY_FRONTEND_DIR = Path(__file__).parent / "frontend" / "dist"
_ASSET_DIR = next((path for path in (_FRONTEND_DIR, _LEGACY_FRONTEND_DIR, _STATIC_DIR) if path.is_dir()), _STATIC_DIR)


@app.get("/", include_in_schema=False)
def frontend():
    return FileResponse(_ASSET_DIR / "index.html")


def _serve_asset(asset_root: Path, asset_path: str):
    asset = (asset_root / asset_path).resolve()
    if asset_root not in asset.parents or not asset.is_file():
        raise HTTPException(status_code=404, detail="Asset not found")
    return FileResponse(asset)


@app.get("/assets/{asset_path:path}", include_in_schema=False)
def frontend_asset(asset_path: str):
    return _serve_asset(_ASSET_DIR, "assets/" + asset_path)


@app.get("/static/{asset_path:path}", include_in_schema=False)
def static_asset(asset_path: str):
    return _serve_asset(_STATIC_DIR, asset_path)


# ---------------------------------------------------------------------------
# Request / Response models
# ---------------------------------------------------------------------------

class GenerateRequest(BaseModel):
    specs: List[str] = Field(default=["easy", "medium", "hard"])
    questions_per_spec: int = Field(default=1, ge=1, le=10)
    excel_id: Optional[str] = None


class StatusUpdateRequest(BaseModel):
    status: str = Field(..., pattern="^(approved|rejected|archived)$")
    review_notes: str = ""


class VerifyRequest(BaseModel):
    ddl_sql: str
    seed_sql: str
    answer_sql: str


class LoginRequest(BaseModel):
    email: EmailStr
    password: str = Field(..., min_length=1, max_length=200)


class RegisterRequest(BaseModel):
    email: EmailStr
    password: str = Field(..., min_length=10, max_length=200)


class GenerationJobRequest(BaseModel):
    specs: List[str] = Field(default=["easy", "medium", "hard"])
    questions_per_spec: int = Field(default=1, ge=1, le=10)
    excel_id: Optional[str] = None
    document_id: Optional[str] = Field(default=None, max_length=100)


class UserUpdateRequest(BaseModel):
    role: Optional[str] = Field(default=None, pattern="^(learner|admin)$")
    status: Optional[str] = Field(default=None, pattern="^(active|locked)$")


class LlmSettingsRequest(BaseModel):
    provider: str = Field(..., pattern="^(anthropic|gemini|nvidia)$")
    model: Optional[str] = Field(default=None, max_length=200)
    base_url: Optional[str] = Field(default=None, max_length=500)
    # Write-only. Never echoed back; omit to keep the stored secret, "" to clear it.
    api_token: Optional[str] = Field(default=None, max_length=500)


# ---------------------------------------------------------------------------
# Background task runner


def _learner_question(question: dict[str, Any]) -> dict[str, Any]:
    return {
        "id": question.get("id"),
        "question_text": question.get("question_text", ""),
        "difficulty": question.get("difficulty"),
        "topic": question.get("topic"),
        "status": question.get("status"),
    }


def _safe_login_response(user: dict[str, Any]) -> dict[str, Any]:
    return {"access_token": create_access_token(user), "token_type": "bearer", "user": public_user(user)}


def _audit(actor: dict[str, Any], action: str, resource_type: str, resource_id: str, summary: str = "") -> None:
    """Record one admin mutation. `summary` must already be redacted."""
    try:
        get_store().save_audit({
            "actor_id": actor.get("id", ""),
            "action": action,
            "resource_type": resource_type,
            "resource_id": resource_id,
            "summary": summary,
        })
    except Exception:
        # Audit storage must never break the mutation it records.
        pass


_SUPPORTED_SUFFIXES = (".csv", ".xlsx", ".json", ".md", ".txt", ".pdf")
_MAX_STORED_TEXT = 20000
_MAX_PREVIEW = 500


class DocumentRejected(ValueError):
    """Import rejection with a message authored here and safe to return."""


def _max_upload_bytes() -> int:
    return int(os.getenv("MAX_UPLOAD_BYTES", "5242880"))


def _document_text(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix in {".md", ".txt", ".json", ".csv"}:
        try:
            return content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise DocumentRejected("File is not valid UTF-8 text") from exc
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        rows = []
        for sheet in workbook.worksheets:
            rows.append(f"# {sheet.title}")
            rows.extend(" | ".join("" if cell is None else str(cell) for cell in row) for row in sheet.iter_rows(values_only=True))
        return "\n".join(rows)
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise DocumentRejected("PDF upload requires pypdf dependency") from exc
        return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(content)).pages)
    raise DocumentRejected(f"Supported files: {', '.join(_SUPPORTED_SUFFIXES)}")


def _seed_questions_from_document(filename: str, text: str) -> list[dict[str, Any]]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".json":
        parsed = json.loads(text)
        items = parsed if isinstance(parsed, list) else parsed.get("questions", [])
        return [item for item in items if isinstance(item, dict) and item.get("question_text")]
    if suffix == ".csv":
        return [{"question_text": row.get("question_text", ""), "difficulty": row.get("difficulty", "medium"), "topic": row.get("topic", "Imported")} for row in csv.DictReader(io.StringIO(text)) if row.get("question_text")]
    return []


# ---------------------------------------------------------------------------
# Background task runner
# ---------------------------------------------------------------------------

# ---------------------------------------------------------------------------
# Background task runner
#
# FastAPI BackgroundTasks runs inside this single API process. It is NOT a
# durable queue: an in-flight job is lost on restart, and nothing retries it.
# Mongo job records are the source of truth for status, so a lost job stays
# visible as `running` rather than disappearing. Production therefore refuses
# to start generation unless ALLOW_BACKGROUND_JOBS is explicitly enabled.
# ---------------------------------------------------------------------------

JOB_STATUSES = ("queued", "running", "succeeded", "failed")

_tasks: dict[str, dict] = {}


def _background_jobs_allowed() -> bool:
    if not is_production():
        return True
    return os.getenv("ALLOW_BACKGROUND_JOBS", "").strip().lower() in {"1", "true", "yes"}


def _job_error_code(exc: Exception) -> str:
    name = exc.__class__.__name__
    if name in {"RuntimeError", "ValueError"} and "LLM" in str(exc).upper():
        return "provider_error"
    return {"RuntimeError": "configuration_error", "ValueError": "invalid_generation_output", "TimeoutError": "provider_timeout"}.get(name, "internal_error")


def _save_job_state(job: dict[str, Any], **changes: Any) -> dict[str, Any]:
    job.update(changes)
    job["updated_at"] = _now_iso()
    _tasks[job["id"]] = job
    try:
        get_store().save_job(dict(job))
    except Exception:
        # Never lose the in-process state because persistence is unavailable.
        pass
    return job


def _persist_generated_questions(question_ids: list[str], job: dict[str, Any], provider: str) -> list[str]:
    """Promote pipeline output into Mongo, the authoritative question store.

    The pipeline still writes each candidate to the legacy SQLite `QuestionBank`
    because the clarity/novelty/semantic/execution validators and the archive
    fallback read from it. That copy is legacy-only: review, listing, and
    learner delivery all read the Mongo record written here.
    """
    qb = QuestionBank()
    store = get_store()
    saved: list[str] = []
    for question_id in question_ids:
        question = qb.get(question_id)
        if not question:
            continue
        record = {
            **question,
            "status": question.get("status") or QuestionStatus.PENDING_REVIEW.value,
            "source": "pipeline",
            "source_job_id": job.get("id", ""),
            "generated_by": job.get("requested_by", ""),
            "provider": provider,
            "legacy_sqlite_id": question_id,
        }
        if job.get("document_id"):
            record["source_document_id"] = job["document_id"]
        if job.get("excel_id"):
            record["source_excel_id"] = job["excel_id"]
        store.save_question(record)
        saved.append(question_id)
    return saved


def _run_generate(
    task_id: str,
    specs: List[str],
    questions_per_spec: int,
    excel_id: Optional[str] = None,
):
    job = _tasks.get(task_id) or get_store().job(task_id) or {"id": task_id}
    overrides = llm_overrides()
    provider = overrides.get("AI_PROVIDER", "")
    try:
        _save_job_state(job, status="running", started_at=_now_iso(), provider=provider)
        result = generate_pipeline(
            specs=specs,
            questions_per_spec=questions_per_spec,
            excel_id=excel_id,
            llm_overrides=overrides,
        )
        saved = _persist_generated_questions(result.get("question_ids", []), job, provider)
        _save_job_state(job, status="succeeded", result=result, persisted_question_ids=saved, finished_at=_now_iso())
        _audit({"id": job.get("requested_by", "")}, "generation.succeeded", "job", task_id, f"questions={len(saved)} provider={provider}")
    except Exception as exc:
        _save_job_state(job, status="failed", error=f"{exc.__class__.__name__}: {exc}", error_code=_job_error_code(exc), finished_at=_now_iso())
        _audit({"id": job.get("requested_by", "")}, "generation.failed", "job", task_id, f"code={_job_error_code(exc)}")


# ---------------------------------------------------------------------------
# Endpoints
# ---------------------------------------------------------------------------

@app.get("/health")
def health():
    try:
        mode = get_store().mode
    except RuntimeError as exc:
        return JSONResponse(status_code=503, content={"status": "unavailable", "version": "2.0.0", "persistence": "unavailable", "detail": str(exc)})
    if mode != "mongo" and is_production():
        return JSONResponse(status_code=503, content={"status": "unavailable", "version": "2.0.0", "persistence": mode, "detail": "Production requires a reachable MongoDB"})
    return {"status": "ok", "version": "2.0.0", "persistence": mode}


@app.post("/auth/register", status_code=201)
def register(req: RegisterRequest):
    store = get_store()
    try:
        user = store.create_user({
            "email": str(req.email),
            "password_hash": hash_password(req.password),
            "role": "learner",
            "status": "active",
        })
    except ValueError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc
    return _safe_login_response(user)


@app.post("/auth/login")
def login(req: LoginRequest):
    user = get_store().find_user_by_email(str(req.email))
    if not user or user.get("status", "active") != "active" or not verify_password(req.password, user.get("password_hash", "")):
        raise HTTPException(status_code=401, detail="Invalid email or password")
    return _safe_login_response(user)


@app.get("/me")
def current_profile(user: dict[str, Any] = Depends(get_current_user)):
    return public_user(user)


@app.post("/generate")
def generate_questions(req: GenerateRequest, background_tasks: BackgroundTasks, _: dict[str, Any] = Depends(require_admin)):
    return create_generation_job(req, background_tasks)


@app.post("/admin/generation-jobs")
def create_generation_job(req: GenerationJobRequest, background_tasks: BackgroundTasks, user: dict[str, Any] = Depends(require_admin)):
    if not _background_jobs_allowed():
        raise HTTPException(
            status_code=503,
            detail="Generation is disabled in production: FastAPI BackgroundTasks is not a durable worker. Set ALLOW_BACKGROUND_JOBS=1 to accept non-durable execution.",
        )
    total = len(req.specs) * req.questions_per_spec
    task_id = str(uuid.uuid4())
    job = {
        "id": task_id,
        "status": "queued",
        "result": None,
        "error": None,
        "error_code": None,
        "requested_by": user["id"],
        "specs": req.specs,
        "questions_per_spec": req.questions_per_spec,
        "excel_id": req.excel_id,
        "document_id": req.document_id,
        "executor": "background_tasks",
        "durable": False,
        "created_at": _now_iso(),
    }
    _save_job_state(job)
    _audit(user, "generation.start", "job", task_id, f"specs={','.join(req.specs)} per_spec={req.questions_per_spec} excel_id={req.excel_id or '-'}")
    background_tasks.add_task(_run_generate, task_id, req.specs, req.questions_per_spec, req.excel_id)
    return JSONResponse(status_code=202, content={"task_id": task_id, "status": "queued", "durable": False, "message": f"Generating {total} questions. Poll GET /admin/generation-jobs/{task_id}"})


def _job_or_404(job_id: str) -> dict[str, Any]:
    """Mongo is the source of truth; process memory only fills a fresher status."""
    try:
        stored = get_store().job(job_id)
    except Exception:
        stored = None
    live = _tasks.get(job_id)
    job = {**(stored or {}), **(live or {})} if (stored or live) else None
    if not job:
        raise HTTPException(status_code=404, detail="Job not found")
    return job


@app.get("/tasks/{task_id}")
def get_task_status(task_id: str, _: dict[str, Any] = Depends(require_admin)):
    return _job_or_404(task_id)


@app.get("/admin/generation-jobs/{job_id}")
def get_generation_job(job_id: str, _: dict[str, Any] = Depends(require_admin)):
    return _job_or_404(job_id)


@app.get("/questions")
def list_questions(
    status: Optional[str] = Query(default=None),
    difficulty: Optional[str] = Query(default=None),
    limit: int = Query(default=50, ge=1, le=200),
    user: dict[str, Any] = Depends(get_current_user),
):
    store = get_store()
    visible_status = status if user.get("role") == "admin" else QuestionStatus.APPROVED.value
    questions = store.questions(status=visible_status, difficulty=difficulty, limit=limit)
    if user.get("role") != "admin":
        questions = [_learner_question(question) for question in questions]
    return {"questions": questions, "total": len(questions)}


@app.get("/questions/{question_id}")
def get_question(question_id: str, user: dict[str, Any] = Depends(get_current_user)):
    q = get_store().question(question_id)
    if q is None:
        raise HTTPException(status_code=404, detail="Question not found")
    if user.get("role") != "admin" and q.get("status") != QuestionStatus.APPROVED.value:
        raise HTTPException(status_code=404, detail="Question not found")
    return q if user.get("role") == "admin" else _learner_question(q)


@app.put("/questions/{question_id}/status")
def update_question_status(question_id: str, req: StatusUpdateRequest, user: dict[str, Any] = Depends(require_admin)):
    store = get_store()
    status_map = {
        "approved": QuestionStatus.APPROVED,
        "rejected": QuestionStatus.REJECTED,
        "archived": QuestionStatus.ARCHIVED,
    }
    new_status = status_map[req.status]
    success = store.update_question_status(question_id, new_status, req.review_notes)
    if not success:
        raise HTTPException(status_code=404, detail="Question not found or update failed")
    _audit(user, "question.review", "question", question_id, f"status={new_status.value if hasattr(new_status, 'value') else new_status}")
    return {"success": True}


@app.get("/stats")
def get_stats(_: dict[str, Any] = Depends(require_admin)):
    return get_store().admin_stats()


@app.get("/admin/stats")
def admin_stats(_: dict[str, Any] = Depends(require_admin)):
    return get_store().admin_stats()


@app.get("/admin/users")
def admin_users(_: dict[str, Any] = Depends(require_admin)):
    return {"users": get_store().users(), "total": len(get_store().users())}


@app.patch("/admin/users/{user_id}")
def update_user(user_id: str, req: UserUpdateRequest, current: dict[str, Any] = Depends(require_admin)):
    if user_id == current["id"] and (req.status == "locked" or req.role == "learner"):
        raise HTTPException(status_code=400, detail="Admin cannot lock or demote self")
    store = get_store()
    user = store.user_by_id(user_id)
    if not user:
        raise HTTPException(status_code=404, detail="User not found")
    if req.role:
        user["role"] = req.role
    if req.status:
        user["status"] = req.status
    saved = store.save_user(user)
    _audit(current, "user.update", "user", user_id, f"role={saved.get('role')} status={saved.get('status')}")
    return public_user(saved)


@app.get("/weekly/current")
def weekly_current(user: dict[str, Any] = Depends(get_current_user)):
    store = get_store()
    challenge = store.current_challenge() or {"id": "none", "number": 42, "title": "Thử thách tuần #42", "status": "active"}
    questions = [_learner_question(q) for q in store.challenge_questions(challenge)]
    done = 0
    total = max(len(questions), 8)
    return {"challenge": challenge, "questions": questions, "progress": {"done": done, "total": total, "percent": round(done / total * 100)}}


@app.get("/admin/documents")
def admin_documents(_: dict[str, Any] = Depends(require_admin)):
    return {"documents": get_store().documents(), "total": len(get_store().documents())}


@app.post("/admin/documents")
async def upload_document(file: UploadFile = File(...), user: dict[str, Any] = Depends(require_admin)):
    filename = file.filename or "upload"
    suffix = Path(filename).suffix.lower()
    if suffix not in _SUPPORTED_SUFFIXES:
        raise HTTPException(status_code=400, detail=f"Supported files: {', '.join(_SUPPORTED_SUFFIXES)}")

    limit = _max_upload_bytes()
    chunks: list[bytes] = []
    size = 0
    while chunk := await file.read(65536):
        size += len(chunk)
        if size > limit:
            raise HTTPException(status_code=413, detail=f"File too large (limit {limit} bytes)")
        chunks.append(chunk)
    content = b"".join(chunks)

    checksum = hashlib.sha256(content).hexdigest()
    store = get_store()
    existing = store.document_by_checksum(checksum)
    if existing is not None:
        raise HTTPException(status_code=409, detail=f"Document already imported (id {existing['id']})")

    try:
        text = _document_text(filename, content)
    except HTTPException:
        raise
    except DocumentRejected as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Could not parse {suffix} file: {exc.__class__.__name__}") from exc

    try:
        document = store.save_document({
            "filename": filename,
            "content_type": file.content_type,
            "checksum": checksum,
            "size": size,
            "text": text[:_MAX_STORED_TEXT],
            "truncated": len(text) > _MAX_STORED_TEXT,
            "preview": text[:_MAX_PREVIEW],
            "uploaded_by": user["id"],
        })
    except ValueError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc
    _audit(user, "document.upload", "document", document["id"], f"{filename} size={size} checksum={checksum[:12]}")
    return document


@app.get("/admin/generation-jobs")
def list_generation_jobs(_: dict[str, Any] = Depends(require_admin)):
    jobs = get_store().jobs()
    return {"jobs": jobs, "total": len(jobs), "durable_executor": False, "statuses": list(JOB_STATUSES)}


@app.get("/admin/audit-logs")
def list_audit_logs(limit: int = Query(default=100, ge=1, le=500), _: dict[str, Any] = Depends(require_admin)):
    logs = get_store().audit_logs(limit=limit)
    return {"logs": logs, "total": len(logs)}


# ---------------------------------------------------------------------------
# Admin LLM settings
#
# The secret lives only in the store/env and is never returned. Responses
# report presence and a short fingerprint, never the token itself.
# ---------------------------------------------------------------------------

_LLM_ENV_KEYS = {
    "anthropic": ("ANTHROPIC_BASE_URL", "ANTHROPIC_MODEL", "ANTHROPIC_AUTH_TOKEN"),
    "gemini": (None, "GEMINI_MODEL", "GEMINI_API_KEY"),
    "nvidia": ("NVIDIA_BASE_URL", "NVIDIA_MODEL", "NVIDIA_API_KEY"),
}


def _stored_llm_settings() -> dict[str, Any]:
    try:
        record = get_store().setting("llm")
    except Exception:
        return {}
    return dict(record.get("value") or {}) if record else {}


def llm_overrides() -> dict[str, str]:
    """Server-side gateway overrides. Secrets stay inside this process."""
    stored = _stored_llm_settings()
    provider = (stored.get("provider") or os.getenv("AI_PROVIDER", "anthropic")).strip().lower()
    base_key, model_key, token_key = _LLM_ENV_KEYS.get(provider, _LLM_ENV_KEYS["anthropic"])
    overrides: dict[str, str] = {"AI_PROVIDER": provider}
    if stored.get("base_url") and base_key:
        overrides[base_key] = stored["base_url"]
    if stored.get("model"):
        overrides[model_key] = stored["model"]
    if stored.get("api_token"):
        overrides[token_key] = stored["api_token"]
    return overrides


def _llm_summary() -> dict[str, Any]:
    stored = _stored_llm_settings()
    provider = (stored.get("provider") or os.getenv("AI_PROVIDER", "anthropic")).strip().lower()
    base_key, model_key, token_key = _LLM_ENV_KEYS.get(provider, _LLM_ENV_KEYS["anthropic"])
    token = stored.get("api_token") or os.getenv(token_key, "")
    model = stored.get("model") or os.getenv(model_key, "")
    base_url = stored.get("base_url") or (os.getenv(base_key, "") if base_key else "")
    return {
        "provider": provider,
        "model": model,
        "base_url": base_url,
        "token_configured": bool(token),
        "token_source": "settings" if stored.get("api_token") else ("env" if token else "none"),
        "token_hint": f"{'*' * 8}{token[-4:]}" if len(token) > 8 else ("*" * len(token) if token else ""),
        "updated_at": stored.get("updated_at", ""),
        "supported_providers": sorted(_LLM_ENV_KEYS),
    }


@app.get("/admin/settings/llm")
def get_llm_settings(_: dict[str, Any] = Depends(require_admin)):
    return _llm_summary()


@app.put("/admin/settings/llm")
def put_llm_settings(req: LlmSettingsRequest, user: dict[str, Any] = Depends(require_admin)):
    stored = _stored_llm_settings()
    value: dict[str, Any] = {
        "provider": req.provider,
        "model": (req.model or stored.get("model") or "").strip(),
        "base_url": (req.base_url or stored.get("base_url") or "").strip(),
        "api_token": stored.get("api_token", ""),
        "updated_at": _now_iso(),
        "updated_by": user["id"],
    }
    if req.api_token is not None:
        # Empty string clears the stored secret and falls back to env.
        value["api_token"] = req.api_token.strip()
    try:
        get_store().save_setting("llm", value)
    except Exception as exc:
        raise HTTPException(status_code=503, detail=f"Settings store unavailable ({exc.__class__.__name__})") from exc
    _audit(user, "settings.llm.update", "settings", "llm", f"provider={value['provider']} model={value['model'] or '-'} token_set={bool(value['api_token'])}")
    return _llm_summary()


@app.post("/admin/llm/test")
def test_llm_provider(_: dict[str, Any] = Depends(require_admin)):
    try:
        gateway = LLMGateway(overrides=llm_overrides())
    except RuntimeError as exc:
        raise HTTPException(status_code=400, detail={"code": "not_configured", "message": str(exc)}) from exc
    try:
        result = gateway.ping()
    except ProviderError as exc:
        raise HTTPException(status_code=502, detail={"code": exc.code, "message": str(exc)}) from exc
    return {"ok": True, **result}


@app.post("/verify-sql")
def verify_sql(req: VerifyRequest, _: dict[str, Any] = Depends(require_admin)):
    try:
        result = verify_in_sandbox(req.ddl_sql, req.seed_sql, req.answer_sql)
        return result
    except Exception as exc:
        raise HTTPException(status_code=500, detail=str(exc))


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    import uvicorn
    uvicorn.run("main:app", host="0.0.0.0", port=8000, reload=True)
